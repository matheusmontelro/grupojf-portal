# app/routers/relatorio_vendas_publico.py

import logging
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import text
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO

from app.database.db import get_db

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/public", tags=["public-reports"])


@router.get("/relatorio-vendas-excel")
async def gerar_relatorio_vendas_excel(
    db: Session = Depends(get_db)
):
    """
    Gera relatório Excel das vendas do período 29/07/2025 a 04/08/2025 com splits.
    Rota pública sem autenticação.
    """
    try:
        # Query para buscar vendas com splits do período específico
        query = text("""
            SELECT 
                -- Dados da Venda
                v.id as venda_id,
                v.documento_venda,
                v.data_venda,
                v.paciente_id,
                p.nome_paciente,
                v.dentista_orcamentista_id,
                d.nome as nome_dentista,
                v.unidade,
                v.valor_bruto_total,
                v.valor_liquido_total,
                v.status_venda,
                v.qtd_parcelas,
                v.pre_venda_id,
                
                -- Dados do Split
                pvs.seller_id,
                pvs.percentage,
                pvs.valor_calculado,
                pvs.absorb_balance,
                pvs.empresa_id,
                pvs.cargo_id,
                pvs.user_id,
                
                -- Dados complementares
                u.nome as nome_beneficiario,
                e.razao_social as nome_empresa,
                c.nome as nome_cargo,
                CASE 
                    WHEN pvs.seller_id = '1' THEN 
                        (SELECT e1.razao_social 
                         FROM empresas e1 
                         WHERE e1.unidade = v.unidade 
                         AND e1.bempaggo_receiver_id = '1' 
                         ORDER BY e1.id 
                         LIMIT 1)
                    ELSE COALESCE(us.nome, es.razao_social)
                END as nome_seller
                
            FROM vendas v
            LEFT JOIN pacientes p ON v.paciente_id = p.paciente_id
            LEFT JOIN dentistas d ON v.dentista_orcamentista_id = d.id
            LEFT JOIN pre_venda_splits pvs ON v.pre_venda_id = pvs.pre_venda_id
            LEFT JOIN users u ON pvs.user_id = u.id
            LEFT JOIN empresas e ON pvs.empresa_id = e.id
            LEFT JOIN cargos c ON pvs.cargo_id = c.id
            LEFT JOIN users us ON pvs.seller_id = us.bempaggo_receiver_id
            LEFT JOIN empresas es ON pvs.seller_id = es.bempaggo_receiver_id
            
            WHERE v.data_venda >= '2025-07-29' 
            AND v.data_venda <= '2025-08-04 23:59:59'
            
            ORDER BY v.data_venda, v.id, pvs.percentage DESC
        """)
        
        result = db.execute(query)
        vendas_splits = result.fetchall()
        
        # Criar workbook Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "Vendas com Splits"
        
        # Estilo para cabeçalho
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # Cabeçalhos
        headers = [
            "ID Venda", "Documento", "Data Venda", "ID Paciente", "Nome Paciente",
            "ID Dentista", "Nome Dentista", "Unidade", "Valor Bruto", "Valor Líquido",
            "Status", "Qtd Parcelas", "Pre-Venda ID", "Seller ID", "Nome Recebedor",
            "Percentual (%)", "Valor Split", "Absorve Saldo", "ID Empresa", "Nome Empresa", 
            "ID Cargo", "Nome Cargo", "ID Usuário", "Nome Beneficiário"
        ]
        
        # Adicionar cabeçalhos
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # Adicionar dados
        for row_idx, row_data in enumerate(vendas_splits, 2):
            ws.cell(row=row_idx, column=1, value=row_data.venda_id)
            ws.cell(row=row_idx, column=2, value=row_data.documento_venda)
            ws.cell(row=row_idx, column=3, value=row_data.data_venda.strftime("%d/%m/%Y %H:%M") if row_data.data_venda else "")
            ws.cell(row=row_idx, column=4, value=row_data.paciente_id)
            ws.cell(row=row_idx, column=5, value=row_data.nome_paciente)
            ws.cell(row=row_idx, column=6, value=row_data.dentista_orcamentista_id)
            ws.cell(row=row_idx, column=7, value=row_data.nome_dentista)
            ws.cell(row=row_idx, column=8, value=row_data.unidade)
            ws.cell(row=row_idx, column=9, value=float(row_data.valor_bruto_total) if row_data.valor_bruto_total else 0)
            ws.cell(row=row_idx, column=10, value=float(row_data.valor_liquido_total) if row_data.valor_liquido_total else 0)
            ws.cell(row=row_idx, column=11, value=row_data.status_venda)
            ws.cell(row=row_idx, column=12, value=row_data.qtd_parcelas)
            ws.cell(row=row_idx, column=13, value=row_data.pre_venda_id)
            ws.cell(row=row_idx, column=14, value=row_data.seller_id)
            ws.cell(row=row_idx, column=15, value=row_data.nome_seller)
            ws.cell(row=row_idx, column=16, value=float(row_data.percentage) * 100 if row_data.percentage else 0)
            ws.cell(row=row_idx, column=17, value=float(row_data.valor_calculado) if row_data.valor_calculado else 0)
            ws.cell(row=row_idx, column=18, value="Sim" if row_data.absorb_balance else "Não" if row_data.absorb_balance is not None else "")
            ws.cell(row=row_idx, column=19, value=row_data.empresa_id)
            ws.cell(row=row_idx, column=20, value=row_data.nome_empresa)
            ws.cell(row=row_idx, column=21, value=row_data.cargo_id)
            ws.cell(row=row_idx, column=22, value=row_data.nome_cargo)
            ws.cell(row=row_idx, column=23, value=row_data.user_id)
            ws.cell(row=row_idx, column=24, value=row_data.nome_beneficiario)
        
        # Formatar colunas de valores como moeda
        for row in range(2, len(vendas_splits) + 2):
            ws.cell(row=row, column=9).number_format = 'R$ #,##0.00'
            ws.cell(row=row, column=10).number_format = 'R$ #,##0.00'
            ws.cell(row=row, column=17).number_format = 'R$ #,##0.00'
            ws.cell(row=row, column=16).number_format = '0,00"%"'
        
        # Ajustar largura das colunas
        column_widths = {
            'A': 10, 'B': 15, 'C': 18, 'D': 12, 'E': 30,
            'F': 12, 'G': 30, 'H': 15, 'I': 15, 'J': 15,
            'K': 12, 'L': 12, 'M': 12, 'N': 10, 'O': 25,
            'P': 12, 'Q': 12, 'R': 12, 'S': 12, 'T': 20, 
            'U': 10, 'V': 20, 'W': 12, 'X': 30
        }
        
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        # Adicionar linha de totais
        total_row = len(vendas_splits) + 3
        ws.cell(row=total_row, column=8, value="TOTAIS:")
        ws.cell(row=total_row, column=8).font = Font(bold=True)
        
        # Fórmulas para totais
        ws.cell(row=total_row, column=9, value=f"=SUM(I2:I{len(vendas_splits)+1})")
        ws.cell(row=total_row, column=10, value=f"=SUM(J2:J{len(vendas_splits)+1})")
        ws.cell(row=total_row, column=17, value=f"=SUM(Q2:Q{len(vendas_splits)+1})")
        
        # Formatar totais
        ws.cell(row=total_row, column=9).number_format = 'R$ #,##0.00'
        ws.cell(row=total_row, column=10).number_format = 'R$ #,##0.00'
        ws.cell(row=total_row, column=17).number_format = 'R$ #,##0.00'
        
        # Adicionar informações do relatório
        info_row = total_row + 2
        ws.cell(row=info_row, column=1, value=f"Relatório gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        ws.cell(row=info_row+1, column=1, value="Período: 29/07/2025 a 04/08/2025")
        ws.cell(row=info_row+2, column=1, value=f"Total de registros: {len(vendas_splits)}")
        
        # Salvar em BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Retornar como download
        filename = f"vendas_splits_29jul_04ago2025_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename={filename}"
            }
        )
        
    except Exception as e:
        logger.error(f"Erro ao gerar relatório Excel: {str(e)}")
        raise HTTPException(
            status_code=500,
            detail=f"Erro ao gerar relatório: {str(e)}"
        )